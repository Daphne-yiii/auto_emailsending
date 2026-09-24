#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import json
import logging
import os
import re
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional
from urllib.parse import urlparse

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from main import (
    ARXIV_DELAY_SECONDS,
    DEFAULT_START_DATE,
    OPENALEX_DELAY_SECONDS,
    REQUEST_TIMEOUT,
    USER_AGENT,
    OpenAlexClient,
    Paper,
    classify,
    choose_email,
    dedupe,
    extract_emails,
    fetch_arxiv,
    normalize_name,
    normalize_space,
    pdf_text,
    plausible_email,
)

DEFAULT_QUERY = (
    'all:"AI for Science" OR all:"AI4Science" OR all:"scientific machine learning" OR '
    'all:"machine learning for science" OR all:"deep learning for science" OR '
    'all:"foundation model" AND (all:biology OR all:chemistry OR all:materials OR '
    'all:physics OR all:protein OR all:molecule OR all:"drug discovery" OR all:climate)'
)
DEFAULT_MAX_PAPERS = 300
DEFAULT_MODEL = os.getenv("OPENAI_MODEL", "gpt-5")
GPT_SEARCH_DELAY_SECONDS = 0.8

OUTPUT_COLUMNS = [
    "姓名",
    "所在机构（学校or哪个公司的哪个部门）",
    "论文链接",
    "论文标题",
    "论文摘要",
    "研究方向标签",
    "关键词",
    "邮箱",
    "GitHub",
    "LinkedIn",
    "个人主页",
    "联系页",
    "邮箱来源",
    "邮箱来源链接",
    "邮箱置信度",
    "华人判断依据",
    "arXiv ID",
    "发表日期",
]

CHINESE_SURNAMES = {
    "ai", "an", "ao", "bai", "bao", "bi", "bian", "cai", "cao", "chang", "che", "chen", "cheng", "chi", "chou",
    "chu", "cui", "dai", "deng", "ding", "dong", "dou", "du", "duan", "fan", "fang", "fei", "feng", "fu",
    "gan", "gao", "ge", "gong", "gu", "guan", "guo", "han", "hao", "he", "hong", "hou", "hu", "hua", "huang",
    "ji", "jia", "jiang", "jin", "kang", "kong", "lai", "lan", "lei", "li", "lian", "liang", "liao", "lin",
    "liu", "long", "lou", "lu", "luo", "ma", "mao", "meng", "mo", "ni", "pan", "pang", "pei", "peng", "qian",
    "qin", "qiu", "qu", "ren", "ruan", "sang", "shao", "shen", "sheng", "shi", "song", "su", "sun", "tan",
    "tang", "tao", "tian", "tong", "wan", "wang", "wei", "wen", "wu", "xia", "xiang", "xiao", "xie", "xin",
    "xing", "xiong", "xu", "xue", "yan", "yang", "yao", "ye", "yi", "yin", "ying", "yu", "yuan", "zeng",
    "zhan", "zhang", "zhao", "zheng", "zhong", "zhou", "zhu", "zhuang", "zou",
}

COMPOUND_CHINESE_SURNAMES = {
    "ouyang", "sima", "situ", "shangguan", "zhuge", "dongfang", "huangfu", "linghu", "murong", "nangong", "ximen",
}


@dataclass
class EmailLookup:
    email: str
    source: str
    source_url: str
    confidence: str
    rationale: str = ""
    github_url: str = ""
    linkedin_url: str = ""
    homepage_url: str = ""
    contact_page_url: str = ""


def normalize_name_token(value: str) -> str:
    value = normalize_name(value)
    value = re.sub(r"[^a-z]", "", value)
    return value


def likely_chinese_name_reason(name: str) -> str:
    tokens = [normalize_name_token(token) for token in re.split(r"[\s\-]+", name or "")]
    tokens = [token for token in tokens if token]
    if not tokens:
        return ""

    compound_candidates = []
    single_candidates = []
    if len(tokens) >= 2:
        compound_candidates.extend([tokens[0] + tokens[1], tokens[-2] + tokens[-1]])
    if tokens:
        single_candidates.extend([tokens[-1], tokens[0]])

    for candidate in compound_candidates + single_candidates:
        if candidate in COMPOUND_CHINESE_SURNAMES:
            return f"中文复姓启发式：{candidate.title()}"
        if candidate in CHINESE_SURNAMES:
            return f"中文姓氏启发式：{candidate.title()}"
    return ""


def is_likely_chinese_name(name: str) -> bool:
    return bool(likely_chinese_name_reason(name))


def email_matches_author(email: str, author_name: str) -> bool:
    if not plausible_email(email):
        return False
    local = email.split("@", 1)[0].lower()
    tokens = [token for token in normalize_name(author_name).split() if len(token) >= 2]
    if not tokens:
        return False
    surname = tokens[-1]
    given = tokens[:-1]
    return (
        surname in local
        or any(token in local for token in given)
        or any(token[0] + surname in local for token in given if token)
        or any(token + surname[0] in local for token in given if surname)
    )


def extract_json_object(text: str) -> dict[str, str]:
    text = (text or "").strip()
    if not text:
        return {}
    try:
        parsed = json.loads(text)
        return parsed if isinstance(parsed, dict) else {}
    except json.JSONDecodeError:
        pass

    match = re.search(r"\{.*\}", text, re.S)
    if not match:
        return {}
    try:
        parsed = json.loads(match.group(0))
        return parsed if isinstance(parsed, dict) else {}
    except json.JSONDecodeError:
        return {}


def normalize_openai_base_url(value: str) -> str:
    value = (value or "").strip().rstrip("/")
    if value and not value.endswith("/v1"):
        value = f"{value}/v1"
    return value


def best_public_contact(results: list[dict[str, object]], author_name: str) -> dict[str, str]:
    emails: list[str] = []
    githubs: list[str] = []
    linkedins: list[str] = []
    homepages: list[str] = []
    contact_pages: list[str] = []

    for result in results:
        link = normalize_space(str(result.get("link", "")))
        title = normalize_space(str(result.get("title", "")))
        snippet = normalize_space(str(result.get("snippet", "")))
        combined = "\n".join([title, snippet, link])
        if not link:
            continue
        emails.extend(extract_emails(combined))
        host = urlparse(link).netloc.lower()
        if "github.com" in host:
            githubs.append(link)
        if "linkedin.com" in host:
            linkedins.append(link)
        if any(token in host for token in [".edu", ".ac.", "github.io", "pages.dev", "vercel.app", "netlify.app"]):
            homepages.append(link)
        if re.search(r"contact|about|people|team|bio|profile", link, re.I) or re.search(r"contact|about|profile|cv|homepage", combined, re.I):
            contact_pages.append(link)

    return {
        "email": choose_email(dedupe(emails), author_name),
        "github": dedupe(githubs)[0] if githubs else "",
        "linkedin": dedupe(linkedins)[0] if linkedins else "",
        "homepage": dedupe(homepages)[0] if homepages else "",
        "contact_page": dedupe(contact_pages)[0] if contact_pages else "",
    }


class GPTWebEmailClient:
    def __init__(self, api_key: str, model: str):
        from openai import OpenAI as _OpenAI

        base_url = normalize_openai_base_url(
            os.getenv("OPENAI_PROXY_BASE_URL", "").strip() or os.getenv("OPENAI_BASE_URL", "").strip()
        )
        kwargs = {"api_key": api_key}
        if base_url:
            kwargs["base_url"] = base_url
        self.client = _OpenAI(**kwargs)
        self.model = model

    def find(self, author_name: str, affiliation: str, paper: Paper) -> EmailLookup:
        prompt = f"""
请联网搜索这位 AI for Science 论文作者的公开联系方式。优先找 GitHub、LinkedIn、个人主页、实验室主页、联系页，以及公开展示的邮箱。

作者姓名：{author_name}
机构：{affiliation or "未知"}
论文标题：{paper.title}
论文链接：{paper.paper_url}

请严格返回一个 JSON 对象，不要输出额外文字：
{{
  "email": "找到的邮箱；没找到则空字符串",
  "source_url": "邮箱来源链接；没找到则空字符串",
  "github_url": "公开 GitHub 主页；没找到则空字符串",
  "linkedin_url": "公开 LinkedIn 主页；没找到则空字符串",
  "homepage_url": "个人主页/实验室主页；没找到则空字符串",
  "contact_page_url": "联系页/关于页；没找到则空字符串",
  "confidence": "high / medium / low / none",
  "rationale": "一句话说明匹配依据"
}}
""".strip()

        try:
            response = self.client.responses.create(
                model=self.model,
                input=prompt,
                tools=[{"type": "web_search"}],
            )
        except Exception as exc:
            try:
                response = self.client.responses.create(
                    model=self.model,
                    input=prompt,
                    tools=[{"type": "web_search_preview"}],
                )
            except Exception as fallback_exc:
                logging.warning("GPT 联网搜索失败：%s | %s | %s", author_name, exc, fallback_exc)
                return EmailLookup("", "gpt_web_search", "", "none", str(fallback_exc))
        finally:
            time.sleep(GPT_SEARCH_DELAY_SECONDS)

        data = extract_json_object(getattr(response, "output_text", ""))
        email = normalize_space(str(data.get("email", ""))).lower()
        email_match = re.search(r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}", email, re.I)
        email = email_match.group(0).lower() if email_match else ""
        confidence = normalize_space(str(data.get("confidence", "none"))).lower()
        if confidence not in {"high", "medium", "low", "none"}:
            confidence = "low" if email else "none"

        source_url = normalize_space(str(data.get("source_url", "")))
        github_url = normalize_space(str(data.get("github_url", "")))
        linkedin_url = normalize_space(str(data.get("linkedin_url", "")))
        homepage_url = normalize_space(str(data.get("homepage_url", "")))
        contact_page_url = normalize_space(str(data.get("contact_page_url", "")))

        if not email or not email_matches_author(email, author_name):
            return EmailLookup(
                "",
                "gpt_web_search",
                source_url,
                "none",
                normalize_space(str(data.get("rationale", ""))),
                github_url,
                linkedin_url,
                homepage_url,
                contact_page_url,
            )

        return EmailLookup(
            email=email,
            source="gpt_web_search",
            source_url=source_url,
            confidence=confidence,
            rationale=normalize_space(str(data.get("rationale", ""))),
            github_url=github_url,
            linkedin_url=linkedin_url,
            homepage_url=homepage_url,
            contact_page_url=contact_page_url,
        )


def pdf_email_lookup(session, paper: Paper, author_name: str, output_dir: Path, no_pdf: bool) -> EmailLookup:
    if no_pdf:
        return EmailLookup("", "pdf_first_5_pages", paper.pdf_url, "none", "PDF fallback disabled")
    text = pdf_text(session, paper, output_dir / "pdf_cache", max_pages=5)
    emails = [email for email in extract_emails(text) if email_matches_author(email, author_name)]
    email = choose_email(emails, author_name)
    if not email:
        return EmailLookup("", "pdf_first_5_pages", paper.pdf_url, "none", "")
    return EmailLookup(email=email, source="pdf_first_5_pages", source_url=paper.pdf_url, confidence="medium", rationale="论文 PDF 前五页提取")


def confidence_rank(value: str) -> int:
    return {"high": 3, "medium": 2, "low": 1, "none": 0}.get((value or "").lower(), 0)


def assign_email(email_to_person: dict[str, tuple[str, int]], email: str, author_name: str, confidence: str) -> bool:
    normalized_author = normalize_name(author_name)
    rank = confidence_rank(confidence)
    existing = email_to_person.get(email)
    if not existing:
        email_to_person[email] = (normalized_author, rank)
        return True
    existing_author, existing_rank = existing
    if existing_author == normalized_author:
        if rank > existing_rank:
            email_to_person[email] = (normalized_author, rank)
        return True
    logging.warning("邮箱冲突，已跳过：%s 已绑定 %s，当前作者 %s", email, existing_author, author_name)
    return False


def build_rows(papers: list[Paper], output_dir: Path, no_pdf: bool, model: str, sleep_seconds: float) -> list[dict[str, str]]:
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("没有检测到 OPENAI_API_KEY；该脚本第一步需要 GPT 联网搜索邮箱。")

    import requests

    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})
    openalex = OpenAlexClient(session, os.getenv("OPENALEX_MAILTO", "").strip())
    gpt_web = GPTWebEmailClient(api_key, model)
    rows: list[dict[str, str]] = []
    email_to_person: dict[str, tuple[str, int]] = {}

    for idx, paper in enumerate(papers, 1):
        logging.info("[%d/%d] %s", idx, len(papers), paper.title[:90])
        chinese_authors = [
            (author, reason)
            for author in paper.authors
            if (reason := likely_chinese_name_reason(author))
        ]
        if not chinese_authors:
            logging.debug("论文没有命中文中姓氏启发式作者，跳过：%s", paper.title[:90])
            continue

        work = openalex.find_work(paper)
        tags, keywords = classify(paper.title, paper.abstract)

        for author, chinese_reason in chinese_authors:
            authorship = openalex.match_author(author, work)
            affiliation = OpenAlexClient.affiliation(authorship)

            logging.info("搜索邮箱：%s | %s", author, affiliation or "未知机构")
            lookup = gpt_web.find(author, affiliation, paper)
            if not lookup.email:
                logging.info("GPT 未找到邮箱，读取 PDF 前五页：%s", author)
                lookup = pdf_email_lookup(session, paper, author, output_dir, no_pdf)

            if not lookup.email:
                logging.info("未找到邮箱，跳过：%s", author)
                continue

            if not assign_email(email_to_person, lookup.email, author, lookup.confidence):
                continue

            rows.append({
                "姓名": author,
                "所在机构（学校or哪个公司的哪个部门）": affiliation or "未找到公开信息",
                "论文链接": paper.paper_url,
                "论文标题": paper.title,
                "论文摘要": paper.abstract,
                "研究方向标签": tags,
                "关键词": keywords,
                "邮箱": lookup.email,
                "GitHub": lookup.github_url,
                "LinkedIn": lookup.linkedin_url,
                "个人主页": lookup.homepage_url,
                "联系页": lookup.contact_page_url,
                "邮箱来源": lookup.source,
                "邮箱来源链接": lookup.source_url,
                "邮箱置信度": lookup.confidence,
                "华人判断依据": chinese_reason,
                "arXiv ID": paper.arxiv_id,
                "发表日期": paper.published.date().isoformat(),
            })
            time.sleep(max(sleep_seconds, 0))
    return rows


def export(rows: list[dict[str, str]], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame(rows, columns=OUTPUT_COLUMNS)
    if not df.empty:
        df = df[df["邮箱"].fillna("").astype(str).str.strip() != ""]
        email_owner_counts = df.groupby("邮箱")["姓名"].nunique()
        conflicting_emails = set(email_owner_counts[email_owner_counts > 1].index)
        if conflicting_emails:
            df = df[~df["邮箱"].isin(conflicting_emails)]
        df = df.drop_duplicates(subset=["姓名", "邮箱", "论文链接"]).reset_index(drop=True)

    xlsx = output_dir / "ai_for_science_chinese_talent_pool.xlsx"
    csv = output_dir / "ai_for_science_chinese_talent_pool.csv"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="人才库")
        ws = writer.book["人才库"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        widths = {
            "A": 22, "B": 45, "C": 36, "D": 60, "E": 100, "F": 35, "G": 45,
            "H": 32, "I": 22, "J": 45, "K": 16, "L": 28, "M": 18, "N": 16,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    df.to_csv(csv, index=False, encoding="utf-8-sig")
    logging.info("已输出：%s", xlsx)
    logging.info("已输出：%s", csv)
    logging.info("有效邮箱人才行数：%d", len(df))


def parse_args() -> argparse.Namespace:
    today = datetime.now(timezone.utc).date().isoformat()
    parser = argparse.ArgumentParser(description="抓取 AI for Science 华人/中国人作者公开邮箱人才库")
    parser.add_argument("--query", default=DEFAULT_QUERY)
    parser.add_argument("--start-date", default=DEFAULT_START_DATE)
    parser.add_argument("--end-date", default=today)
    parser.add_argument("--max-papers", type=int, default=DEFAULT_MAX_PAPERS)
    parser.add_argument("--output-dir", default="output")
    parser.add_argument("--model", default=DEFAULT_MODEL)
    parser.add_argument("--sleep", type=float, default=0.2)
    parser.add_argument("--no-pdf", action="store_true")
    parser.add_argument("--verbose", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")
    if "replace-with-your-email" in USER_AGENT:
        logging.warning("请先把 main.py 里的 USER_AGENT 联系邮箱替换成你自己的邮箱。")

    papers = fetch_arxiv(args.query, args.start_date, args.end_date, args.max_papers)
    logging.info("共获取 %d 篇 AI for Science 候选论文", len(papers))
    if not papers:
        return 0

    output_dir = Path(args.output_dir)
    rows = build_rows(papers, output_dir, args.no_pdf, args.model, args.sleep)
    export(rows, output_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
