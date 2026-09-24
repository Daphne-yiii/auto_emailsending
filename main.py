#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import argparse
import html
import logging
import os
import re
import time
import unicodedata
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Optional
from urllib.parse import urlparse

import arxiv
import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader
from rapidfuzz import fuzz

DEFAULT_QUERY = 'all:"multi-agent" OR all:"multi agent" OR all:multiagent'
DEFAULT_START_DATE = "2024-01-01"
DEFAULT_MAX_PAPERS = 300
REQUEST_TIMEOUT = 30
ARXIV_DELAY_SECONDS = 3.2
OPENALEX_DELAY_SECONDS = 0.2
GITHUB_DELAY_SECONDS = 0.3
SEARCH_DELAY_SECONDS = 0.6
USER_AGENT = "multi-agent-talent-research/1.0 (contact: replace-with-your-email@example.com)"

OUTPUT_COLUMNS = [
    "姓名",
    "所在机构（学校or哪个公司的哪个部门）",
    "论文链接",
    "论文标题",
    "论文摘要",
    "研究方向标签",
    "关键词",
    "联系方式类型",
    "联系方式内容",
]

STANDARD_EMAIL_RE = re.compile(r"(?<![\w.+-])([A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,})(?![\w.-])", re.I)
OBFUSCATED_EMAIL_RE = re.compile(
    r"(?<![\w.+-])([A-Z0-9._%+\-]+(?:\s*(?:\[at\]|\(at\)| at |@)\s*)"
    r"[A-Z0-9.\-]+(?:\s*(?:\[dot\]|\(dot\)| dot |\.)\s*)[A-Z]{2,})(?![\w.-])",
    re.I,
)
GENERIC_PREFIXES = {"info", "contact", "admin", "support", "help", "office", "team", "press", "media", "sales", "hello", "noreply", "no-reply"}
BAD_DOMAINS = {"example.com", "email.com", "domain.com", "test.com"}

RESEARCH_TAG_RULES: dict[str, list[str]] = {
    "多智能体系统": ["multi-agent", "multi agent", "multiagent", "agent collaboration", "agent cooperation", "agent coordination"],
    "强化学习": ["reinforcement learning", "multi-agent reinforcement learning", "marl", "policy gradient", "q-learning", "actor-critic"],
    "大语言模型": ["large language model", "llm", "language model", "foundation model", "prompting", "reasoning"],
    "迁移学习": ["transfer learning", "domain adaptation", "domain generalization", "few-shot", "zero-shot", "meta-learning"],
    "计算机视觉": ["computer vision", "visual", "image", "video", "vision-language", "object detection", "segmentation", "tracking"],
    "机器人": ["robot", "robotics", "embodied", "navigation", "manipulation", "swarm robotics", "autonomous vehicle", "drone", "uav"],
    "自然语言处理": ["natural language processing", "nlp", "text generation", "dialogue", "question answering", "information extraction"],
    "博弈论与机制设计": ["game theory", "mechanism design", "auction", "equilibrium", "negotiation", "social dilemma"],
    "规划与搜索": ["planning", "task planning", "tree search", "monte carlo tree search", "mcts", "path planning"],
    "知识图谱": ["knowledge graph", "graph neural network", "gnn", "graph learning", "relational reasoning"],
    "数据科学": ["data science", "data mining", "data analysis", "analytics", "dataset", "benchmark"],
    "安全与可信AI": ["safety", "alignment", "robustness", "adversarial", "privacy", "trustworthy", "security", "hallucination"],
    "推荐系统": ["recommender", "recommendation system", "collaborative filtering", "ranking"],
    "分布式系统": ["distributed system", "decentralized", "federated learning", "consensus", "communication-efficient"],
}

@dataclass
class Paper:
    arxiv_id: str
    title: str
    abstract: str
    paper_url: str
    pdf_url: str
    published: datetime
    authors: list[str]


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", html.unescape(value or "")).strip()


def normalize_name(value: str) -> str:
    value = unicodedata.normalize("NFKD", value or "")
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = re.sub(r"[^\w\s\-'.]", " ", value)
    return normalize_space(value).lower()


def dedupe(values: Iterable[str]) -> list[str]:
    seen, result = set(), []
    for value in values:
        value = normalize_space(value)
        key = value.lower()
        if value and key not in seen:
            seen.add(key)
            result.append(value)
    return result


def request(session: requests.Session, method: str, url: str, **kwargs: Any) -> Optional[requests.Response]:
    for attempt in range(1, 5):
        try:
            resp = session.request(method, url, timeout=REQUEST_TIMEOUT, **kwargs)
            if resp.status_code == 429:
                time.sleep(min(60, 2 ** attempt * 2))
                continue
            resp.raise_for_status()
            return resp
        except requests.RequestException as exc:
            if attempt == 4:
                logging.warning("请求失败：%s | %s", url, exc)
                return None
            time.sleep(min(30, 2 ** attempt))
    return None


def normalize_email(raw: str) -> str:
    raw = re.sub(r"\s*(?:\[at\]|\(at\)|\bat\b)\s*", "@", raw, flags=re.I)
    raw = re.sub(r"\s*(?:\[dot\]|\(dot\)|\bdot\b)\s*", ".", raw, flags=re.I)
    return re.sub(r"\s+", "", raw).strip(".,;:()[]{}<>").lower()


def plausible_email(email: str) -> bool:
    if not STANDARD_EMAIL_RE.fullmatch(email):
        return False
    prefix, domain = email.split("@", 1)
    return prefix not in GENERIC_PREFIXES and domain not in BAD_DOMAINS


def extract_emails(text: str) -> list[str]:
    out = []
    for m in OBFUSCATED_EMAIL_RE.finditer(text or ""):
        email = normalize_email(m.group(1))
        if plausible_email(email):
            out.append(email)
    return dedupe(out)


def choose_email(emails: list[str], author_name: str) -> str:
    tokens = [t for t in normalize_name(author_name).split() if len(t) >= 2]
    scored = []
    for email in emails:
        local = email.split("@", 1)[0]
        score = sum(3 if token in local else 1 if token[0] in local else 0 for token in tokens)
        if ".edu" in email or ".ac." in email:
            score += 1
        scored.append((score, email))
    scored.sort(key=lambda x: (-x[0], x[1]))
    return scored[0][1] if scored and scored[0][0] > 0 else ""


def fetch_arxiv(query: str, start_date: str, end_date: str, max_papers: int) -> list[Paper]:
    start_dt = datetime.fromisoformat(start_date).replace(tzinfo=timezone.utc)
    end_dt = datetime.fromisoformat(end_date).replace(hour=23, minute=59, second=59, tzinfo=timezone.utc)
    client = arxiv.Client(page_size=min(100, max_papers), delay_seconds=ARXIV_DELAY_SECONDS, num_retries=4)
    search = arxiv.Search(query=query, max_results=max_papers * 3, sort_by=arxiv.SortCriterion.SubmittedDate, sort_order=arxiv.SortOrder.Descending)
    papers, seen = [], set()
    for result in client.results(search):
        published = result.published.astimezone(timezone.utc)
        if published > end_dt:
            continue
        if published < start_dt:
            break
        arxiv_id = result.get_short_id().split("v")[0]
        if arxiv_id in seen:
            continue
        seen.add(arxiv_id)
        papers.append(Paper(
            arxiv_id=arxiv_id,
            title=normalize_space(result.title),
            abstract=normalize_space(result.summary),
            paper_url=result.entry_id,
            pdf_url=result.pdf_url,
            published=published,
            authors=[normalize_space(a.name) for a in result.authors],
        ))
        if len(papers) >= max_papers:
            break
    return papers


class OpenAlexClient:
    def __init__(self, session: requests.Session, mailto: str = ""):
        self.session, self.mailto = session, mailto

    def find_work(self, paper: Paper) -> Optional[dict[str, Any]]:
        params = {"search": paper.title, "per-page": 5, "select": "id,display_name,authorships,topics,keywords"}
        if self.mailto:
            params["mailto"] = self.mailto
        resp = request(self.session, "GET", "https://api.openalex.org/works", params=params)
        time.sleep(OPENALEX_DELAY_SECONDS)
        if not resp:
            return None
        best_score, best = 0, None
        for work in resp.json().get("results", []):
            score = fuzz.token_set_ratio(normalize_name(paper.title), normalize_name(work.get("display_name", "")))
            if score > best_score:
                best_score, best = score, work
        return best if best_score >= 88 else None

    @staticmethod
    def match_author(author_name: str, work: Optional[dict[str, Any]]) -> Optional[dict[str, Any]]:
        if not work:
            return None
        best_score, best = 0, None
        for authorship in work.get("authorships", []):
            candidate = authorship.get("author", {}).get("display_name", "")
            score = fuzz.token_set_ratio(normalize_name(author_name), normalize_name(candidate))
            if score > best_score:
                best_score, best = score, authorship
        return best if best_score >= 85 else None

    @staticmethod
    def affiliation(authorship: Optional[dict[str, Any]]) -> str:
        if not authorship:
            return ""
        names = [normalize_space(x.get("display_name", "")) for x in authorship.get("institutions", []) if x.get("display_name")]
        return "；".join(dedupe(names))


def pdf_text(session: requests.Session, paper: Paper, cache_dir: Path, max_pages: int = 3) -> str:
    cache_dir.mkdir(parents=True, exist_ok=True)
    path = cache_dir / f"{paper.arxiv_id.replace('/', '_')}.pdf"
    if not path.exists():
        resp = request(session, "GET", paper.pdf_url, headers={"User-Agent": USER_AGENT})
        if not resp:
            return ""
        path.write_bytes(resp.content)
        time.sleep(ARXIV_DELAY_SECONDS)
    try:
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages[:max_pages])
    except Exception:
        return ""


def infer_affiliation_from_pdf(text: str) -> str:
    lines = [normalize_space(x) for x in (text or "").splitlines()[:120]]
    pattern = re.compile(r"\b(university|college|institute|laboratory|lab\b|research center|research centre|department|school of|faculty of|academy|corporation|inc\.?|ltd\.?|technologies|research)\b", re.I)
    candidates = [x for x in lines if 5 <= len(x) <= 220 and pattern.search(x) and "arxiv" not in x.lower()]
    return min(candidates, key=len) if candidates else ""


class GitHubClient:
    def __init__(self, session: requests.Session, token: str = ""):
        self.session = session
        self.headers = {"Accept": "application/vnd.github+json", "User-Agent": USER_AGENT, "X-GitHub-Api-Version": "2022-11-28"}
        if token:
            self.headers["Authorization"] = f"Bearer {token}"

    def find(self, author_name: str, affiliation: str) -> dict[str, str]:
        resp = request(self.session, "GET", "https://api.github.com/search/users", headers=self.headers, params={"q": f'"{author_name}" in:fullname', "per_page": 10})
        time.sleep(GITHUB_DELAY_SECONDS)
        if not resp:
            return {}
        best_score, best = -1, {}
        for item in resp.json().get("items", []):
            login = item.get("login")
            if not login:
                continue
            profile_resp = request(self.session, "GET", f"https://api.github.com/users/{login}", headers=self.headers)
            time.sleep(GITHUB_DELAY_SECONDS)
            if not profile_resp:
                continue
            profile = profile_resp.json()
            score = fuzz.token_set_ratio(normalize_name(author_name), normalize_name(profile.get("name", "")))
            aff = normalize_name(affiliation)
            company = normalize_name(profile.get("company", ""))
            bio = normalize_name(profile.get("bio", ""))
            if aff and (fuzz.partial_ratio(aff, company) > 65 or fuzz.partial_ratio(aff, bio) > 65):
                score += 20
            if score > best_score:
                best_score, best = score, profile
        if best_score < 82 or not best:
            return {}
        email = normalize_space(best.get("email") or "")
        if email and not plausible_email(email):
            email = ""
        return {"github": normalize_space(best.get("html_url") or ""), "email": email, "company": normalize_space(best.get("company") or "")}


class SerperClient:
    def __init__(self, session: requests.Session, api_key: str):
        self.session, self.api_key = session, api_key

    def search(self, q: str) -> list[dict[str, Any]]:
        resp = request(self.session, "POST", "https://google.serper.dev/search", headers={"X-API-KEY": self.api_key, "Content-Type": "application/json"}, json={"q": q, "num": 8})
        time.sleep(SEARCH_DELAY_SECONDS)
        return resp.json().get("organic", []) if resp else []

    def find(self, author_name: str, affiliation: str, title: str) -> dict[str, str]:
        queries = [
            f'"{author_name}" "{affiliation.split("；")[0] if affiliation else ""}" email',
            f'"{author_name}" "{title}"',
            f'"{author_name}" site:github.com',
            f'"{author_name}" site:linkedin.com/in',
        ]
        results = []
        for q in queries:
            results.extend(self.search(q))
        emails, githubs, linkedins, wechats = [], [], [], []
        for result in results:
            link = normalize_space(result.get("link", ""))
            combined = "\n".join([normalize_space(result.get("title", "")), normalize_space(result.get("snippet", "")), link])
            emails.extend(extract_emails(combined))
            host = urlparse(link).netloc.lower()
            if "github.com" in host:
                githubs.append(link)
            if "linkedin.com" in host and "/in/" in link:
                linkedins.append(link)
            m = re.search(r"(?:微信|WeChat)\s*(?:ID|号|：|:)?\s*([A-Za-z][A-Za-z0-9_-]{5,19})", combined, re.I)
            if m:
                wechats.append(m.group(1))
        return {
            "email": choose_email(dedupe(emails), author_name),
            "github": dedupe(githubs)[0] if githubs else "",
            "linkedin": dedupe(linkedins)[0] if linkedins else "",
            "wechat": dedupe(wechats)[0] if wechats else "",
        }


def classify(title: str, abstract: str) -> tuple[str, str]:
    text = f"{title}\n{abstract}".lower()
    tags, keywords = [], []
    for tag, words in RESEARCH_TAG_RULES.items():
        score = sum(1 for word in words if word in text)
        if score:
            tags.append((score, tag))
            keywords.extend(word for word in words if word in text)
    tags.sort(key=lambda x: (-x[0], x[1]))
    keywords = sorted(dedupe(keywords), key=lambda x: (-text.count(x.lower()), -len(x), x))
    return "；".join([x[1] for x in tags[:5]] or ["多智能体系统"]), "；".join(keywords[:12] or ["multi-agent"])


def select_contact(email: str, wechat: str, linkedin: str, github: str) -> tuple[str, str]:
    if email:
        return "邮箱", email
    if wechat:
        return "微信", wechat
    if linkedin:
        return "领英", linkedin
    if github:
        return "GitHub", github
    return "未找到公开联系方式", ""


def build_rows(papers: list[Paper], output_dir: Path, no_pdf: bool) -> list[dict[str, str]]:
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})
    openalex = OpenAlexClient(session, os.getenv("OPENALEX_MAILTO", "").strip())
    github = GitHubClient(session, os.getenv("GITHUB_TOKEN", "").strip())
    serper_key = os.getenv("SERPER_API_KEY", "").strip()
    serper = SerperClient(session, serper_key) if serper_key else None
    rows = []

    for idx, paper in enumerate(papers, 1):
        logging.info("[%d/%d] %s", idx, len(papers), paper.title[:90])
        work = openalex.find_work(paper)
        text = "" if no_pdf else pdf_text(session, paper, output_dir / "pdf_cache")
        emails = extract_emails(text)
        tags, keywords = classify(paper.title, paper.abstract)

        for author in paper.authors:
            authorship = openalex.match_author(author, work)
            affiliation = openalex.affiliation(authorship)
            if not affiliation and text:
                affiliation = infer_affiliation_from_pdf(text)

            email = choose_email(emails, author)
            github_url = linkedin_url = wechat = ""

            if not email:
                gh = github.find(author, affiliation)
                email = gh.get("email", "")
                github_url = gh.get("github", "")
                if not affiliation:
                    affiliation = gh.get("company", "")

            if serper and (not email or not linkedin_url):
                web = serper.find(author, affiliation, paper.title)
                email = email or web.get("email", "")
                github_url = github_url or web.get("github", "")
                linkedin_url = web.get("linkedin", "")
                wechat = web.get("wechat", "")

            contact_type, contact_value = select_contact(email, wechat, linkedin_url, github_url)
            rows.append({
                "姓名": author,
                "所在机构（学校or哪个公司的哪个部门）": affiliation or "未找到公开信息",
                "论文链接": paper.paper_url,
                "论文标题": paper.title,
                "论文摘要": paper.abstract,
                "研究方向标签": tags,
                "关键词": keywords,
                "联系方式类型": contact_type,
                "联系方式内容": contact_value,
            })
    return rows


def export(rows: list[dict[str, str]], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame(rows, columns=OUTPUT_COLUMNS).drop_duplicates(subset=["姓名", "论文链接"]).reset_index(drop=True)
    xlsx = output_dir / "multi_agent_arxiv_talent_pool.xlsx"
    csv = output_dir / "multi_agent_arxiv_talent_pool.csv"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="人才库")
        ws = writer.book["人才库"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        widths = {"A": 22, "B": 45, "C": 36, "D": 60, "E": 100, "F": 35, "G": 45, "H": 20, "I": 45}
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    df.to_csv(csv, index=False, encoding="utf-8-sig")
    logging.info("已输出：%s", xlsx)
    logging.info("已输出：%s", csv)


def main() -> None:
    today = datetime.now(timezone.utc).date().isoformat()
    parser = argparse.ArgumentParser()
    parser.add_argument("--query", default=DEFAULT_QUERY)
    parser.add_argument("--start-date", default=DEFAULT_START_DATE)
    parser.add_argument("--end-date", default=today)
    parser.add_argument("--max-papers", type=int, default=DEFAULT_MAX_PAPERS)
    parser.add_argument("--output-dir", default="output")
    parser.add_argument("--no-pdf", action="store_true")
    parser.add_argument("--verbose", action="store_true")
    args = parser.parse_args()

    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")
    if "replace-with-your-email" in USER_AGENT:
        logging.warning("请先把 USER_AGENT 里的联系邮箱替换成你自己的邮箱。")

    papers = fetch_arxiv(args.query, args.start_date, args.end_date, args.max_papers)
    logging.info("共获取 %d 篇论文", len(papers))
    if not papers:
        return
    output_dir = Path(args.output_dir)
    rows = build_rows(papers, output_dir, args.no_pdf)
    export(rows, output_dir)


if __name__ == "__main__":
    main()
